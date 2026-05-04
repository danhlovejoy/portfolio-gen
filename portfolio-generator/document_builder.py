"""DOCX assignment sheet and rubric builder using python-docx."""

import os
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from rubric_generator import NLP_CLOS, CV_CLOS, CRITERIA

# ---------------------------------------------------------------------------
# Colors
# ---------------------------------------------------------------------------

PRIMARY = RGBColor(26, 58, 92)       # #1a3a5c
ACCENT = RGBColor(46, 125, 50)       # #2e7d32
MUTED = RGBColor(100, 100, 100)
WHITE = RGBColor(255, 255, 255)
HEADER_FILL = "1A3A5C"
LIGHT_FILL = "F7F9FC"
WHITE_FILL = "FFFFFF"


def build_docx(student_name, presentation_type, selected_nlp_indices, selected_cv_indices,
               project_description, contingency, rubric_data, output_dir="/tmp"):
    """Build a DOCX with the assignment sheet and customized rubric.

    Returns:
        str: path to the generated DOCX file.
    """
    # Determine course label
    if presentation_type == "Combined (NLP + CV)":
        course_label = "AIML 2003 (NLP) + AIML 2013 (CV)"
    elif presentation_type == "Standalone NLP":
        course_label = "AIML 2003: Introduction to Natural Language Processing"
    else:
        course_label = "AIML 2013: Introduction to Computer Vision"

    doc = Document()

    # ----- Page setup (portrait, US Letter) -----
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # ----- Header -----
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run(f"Rose State College  |  {course_label}")
    hr.font.size = Pt(9)
    hr.font.bold = True
    hr.font.color.rgb = MUTED

    # ----- Footer -----
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_str = datetime.now().strftime("%B %d, %Y")
    fr = fp.add_run(f"Generated {date_str}  |  Instructor: Dan Lovejoy")
    fr.font.size = Pt(8)
    fr.font.italic = True
    fr.font.color.rgb = RGBColor(150, 150, 150)

    # ===== TITLE =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.space_after = Pt(2)
    tr = title.add_run("Final Portfolio Presentation")
    tr.font.size = Pt(20)
    tr.font.bold = True
    tr.font.color.rgb = PRIMARY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.space_after = Pt(12)
    sr = subtitle.add_run("Custom Assignment Sheet")
    sr.font.size = Pt(12)
    sr.font.color.rgb = MUTED

    # ===== STUDENT INFO TABLE =====
    info_table = doc.add_table(rows=2, cols=4)
    info_table.autofit = True
    _shade_table(info_table, LIGHT_FILL)
    _set_table_borders(info_table, "D0D7DE")

    _info_cell(info_table.cell(0, 0), "Student:", bold=True)
    _info_cell(info_table.cell(0, 1), student_name)
    _info_cell(info_table.cell(0, 2), "Type:", bold=True)
    _info_cell(info_table.cell(0, 3), presentation_type)
    _info_cell(info_table.cell(1, 0), "Date:", bold=True)
    _info_cell(info_table.cell(1, 1), "Tuesday, May 12, 2026")
    _info_cell(info_table.cell(1, 2), "Duration:", bold=True)
    _info_cell(info_table.cell(1, 3), "8–10 minutes")

    doc.add_paragraph()  # spacer

    # ===== SELECTED CLOs =====
    _section_heading(doc, "Selected Course Learning Outcomes")

    if selected_nlp_indices:
        if presentation_type == "Combined (NLP + CV)":
            sh = doc.add_paragraph()
            r = sh.add_run("AIML 2003 — Natural Language Processing")
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = PRIMARY
        for i in selected_nlp_indices:
            p = doc.add_paragraph(f"{i + 1}.  {NLP_CLOS[i]}")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25)
            for run in p.runs:
                run.font.size = Pt(10)

    if selected_cv_indices:
        if presentation_type == "Combined (NLP + CV)":
            sh = doc.add_paragraph()
            r = sh.add_run("AIML 2013 — Computer Vision")
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = PRIMARY
        for i in selected_cv_indices:
            p = doc.add_paragraph(f"{i + 1}.  {CV_CLOS[i]}")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25)
            for run in p.runs:
                run.font.size = Pt(10)

    # ===== PROJECT DESCRIPTION =====
    _section_heading(doc, "Project Description")
    p = doc.add_paragraph(project_description)
    for run in p.runs:
        run.font.size = Pt(10)

    # ===== CONTINGENCY PLAN =====
    _section_heading(doc, "Contingency Plan")
    p = doc.add_paragraph(contingency)
    for run in p.runs:
        run.font.size = Pt(10)

    # ===== SIGNATURE LINES =====
    doc.add_paragraph()  # spacer
    sig = doc.add_paragraph()
    sig.space_before = Pt(24)
    sig_run = sig.add_run("________________________________________          ____________________")
    sig_run.font.size = Pt(10)
    sig_run.font.color.rgb = RGBColor(180, 180, 180)

    sig_label = doc.add_paragraph()
    r1 = sig_label.add_run("Instructor approval")
    r1.font.size = Pt(9)
    r1.font.italic = True
    r1.font.color.rgb = MUTED
    r2 = sig_label.add_run("                                                        Date")
    r2.font.size = Pt(9)
    r2.font.italic = True
    r2.font.color.rgb = MUTED

    # ===== RUBRIC PAGE (landscape) =====
    new_section = doc.add_section()
    new_section.orientation = WD_ORIENT.LANDSCAPE
    new_section.page_width = Inches(11)
    new_section.page_height = Inches(8.5)
    new_section.top_margin = Inches(0.75)
    new_section.bottom_margin = Inches(0.75)
    new_section.left_margin = Inches(0.75)
    new_section.right_margin = Inches(0.75)

    # Copy header/footer to new section
    new_header = new_section.header
    new_header.is_linked_to_previous = False
    nhp = new_header.paragraphs[0]
    nhp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nhr = nhp.add_run(f"Rose State College  |  {course_label}")
    nhr.font.size = Pt(9)
    nhr.font.bold = True
    nhr.font.color.rgb = MUTED

    new_footer = new_section.footer
    new_footer.is_linked_to_previous = False
    nfp = new_footer.paragraphs[0]
    nfp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nfr = nfp.add_run(f"Generated {date_str}  |  Instructor: Dan Lovejoy")
    nfr.font.size = Pt(8)
    nfr.font.italic = True
    nfr.font.color.rgb = RGBColor(150, 150, 150)

    # Rubric title
    rt = doc.add_paragraph()
    rt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt.space_after = Pt(2)
    rtr = rt.add_run(f"Rubric: {student_name}")
    rtr.font.size = Pt(16)
    rtr.font.bold = True
    rtr.font.color.rgb = PRIMARY

    rs = doc.add_paragraph()
    rs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs.space_after = Pt(12)
    rsr = rs.add_run("100 points total  |  Each criterion scored Full / Partial / Minimal / No Credit")
    rsr.font.size = Pt(10)
    rsr.font.color.rgb = MUTED

    # ===== RUBRIC TABLE =====
    _build_rubric_table(doc, rubric_data)

    # ----- Save -----
    safe_name = "".join(
        c if c.isalnum() or c in " -_" else "" for c in student_name
    ).strip().replace(" ", "_")
    filename = f"portfolio_assignment_{safe_name}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _section_heading(doc, text):
    """Add a section heading with accent-colored bottom border."""
    p = doc.add_paragraph()
    p.space_before = Pt(16)
    p.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = PRIMARY
    # Add bottom border
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="1" w:color="2E7D32"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


def _info_cell(cell, text, bold=False):
    """Set text in an info-table cell."""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.space_before = Pt(4)
    p.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(10)
    if bold:
        r.font.bold = True
        r.font.color.rgb = PRIMARY


def _shade_cell(cell, color_hex):
    """Apply shading to a table cell."""
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading)


def _shade_table(table, color_hex):
    """Apply shading to all cells in a table."""
    for row in table.rows:
        for cell in row.cells:
            _shade_cell(cell, color_hex)


def _set_table_borders(table, color_hex):
    """Set thin borders on all sides of a table."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def _build_rubric_table(doc, rubric_data):
    """Build the 4-criteria x 4-level rubric table."""

    levels = ("full", "partial", "minimal", "none")
    level_labels = ("Full Credit", "Partial Credit", "Minimal Credit", "No Credit")

    # 5 columns: criterion + 4 levels
    table = doc.add_table(rows=1 + len(CRITERIA), cols=5)
    table.autofit = True
    _set_table_borders(table, "D0D7DE")

    # ----- Header row -----
    header_texts = ["Criterion"] + list(level_labels)
    for i, text in enumerate(header_texts):
        cell = table.rows[0].cells[i]
        _shade_cell(cell, HEADER_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.space_before = Pt(6)
        p.space_after = Pt(6)
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = WHITE

    # ----- Data rows -----
    for row_idx, criterion in enumerate(CRITERIA):
        name = criterion["name"]
        points = criterion["points"]
        descs = rubric_data.get(name, {})
        row = table.rows[row_idx + 1]
        fill = LIGHT_FILL if row_idx % 2 == 0 else WHITE_FILL

        # Criterion name cell
        cell0 = row.cells[0]
        _shade_cell(cell0, fill)
        cell0.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        p0 = cell0.paragraphs[0]
        p0.space_before = Pt(6)
        p0.space_after = Pt(6)
        r_name = p0.add_run(name)
        r_name.font.size = Pt(9)
        r_name.font.bold = True
        r_name.font.color.rgb = PRIMARY
        p0.add_run("\n")
        r_pts = p0.add_run(f"({points} pts)")
        r_pts.font.size = Pt(8)
        r_pts.font.color.rgb = MUTED

        # Level cells
        for col_idx, level in enumerate(levels):
            cell = row.cells[col_idx + 1]
            _shade_cell(cell, fill)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            p.space_before = Pt(4)
            p.space_after = Pt(4)
            r = p.add_run(descs.get(level, "—"))
            r.font.size = Pt(8)
